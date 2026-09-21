import hashlib
import io
import tempfile
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

import tiktoken
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from docx import Document as DocxDocument
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from .models import Document, DocumentChunk
from .services.document_processing import (
    TextBlock,
    create_chunks,
    extract_document_text,
    normalize_text,
)
from .tasks import process_document


PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def txt_upload(content: bytes = b"Guvenli dokuman metni", name: str = "guide.txt"):
    return SimpleUploadedFile(name, content, content_type="text/plain")


def text_pdf_upload():
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 760, "Birinci sayfa uretim proseduru")
    pdf.showPage()
    pdf.drawString(72, 760, "Ikinci sayfa kalite kontrolu")
    pdf.save()
    return SimpleUploadedFile("guide.pdf", output.getvalue(), content_type=PDF_MIME)


def blank_pdf_upload():
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(output)
    return SimpleUploadedFile("scan.pdf", output.getvalue(), content_type=PDF_MIME)


def encrypted_pdf_upload():
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    writer.write(output)
    return SimpleUploadedFile("secret.pdf", output.getvalue(), content_type=PDF_MIME)


def docx_upload():
    output = io.BytesIO()
    document = DocxDocument()
    document.add_heading("Kalite", level=1)
    document.add_paragraph("Parca kontrol adimlari burada aciklanir.")
    document.save(output)
    return SimpleUploadedFile("guide.docx", output.getvalue(), content_type=DOCX_MIME)


def macro_docx_upload():
    original = docx_upload()
    source = io.BytesIO(original.read())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_archive:
        with zipfile.ZipFile(output, "w") as target_archive:
            for member in source_archive.infolist():
                target_archive.writestr(member, source_archive.read(member.filename))
            target_archive.writestr("word/vbaProject.bin", b"not-a-real-macro")
    return SimpleUploadedFile("macro.docx", output.getvalue(), content_type=DOCX_MIME)


class ProcessingTestBase(TestCase):
    def setUp(self):
        super().setUp()
        self.media_directory = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.media_directory.name)
        self.settings_override.enable()
        self.staff = get_user_model().objects.create(username="processing-staff", is_staff=True)

    def tearDown(self):
        self.settings_override.disable()
        self.media_directory.cleanup()
        super().tearDown()

    def create_document(self, upload=None, title="Prosedür"):
        return Document.objects.create(
            title=title,
            file=upload or txt_upload(),
            uploaded_by=self.staff,
        )


class TextExtractionTests(ProcessingTestBase):
    def test_pdf_text_and_page_numbers_are_extracted(self):
        document = self.create_document(text_pdf_upload())

        blocks = extract_document_text(document)

        self.assertEqual([block.page_number for block in blocks], [1, 2])
        self.assertIn("Birinci sayfa", blocks[0].content)
        self.assertIn("Ikinci sayfa", blocks[1].content)

    def test_docx_text_and_heading_are_extracted(self):
        document = self.create_document(docx_upload())

        blocks = extract_document_text(document)

        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0].section_title, "Kalite")
        self.assertIn("kontrol adimlari", blocks[0].content)

    def test_utf8_txt_is_extracted_and_normalized(self):
        document = self.create_document(
            txt_upload("  U\u0308retim\r\n\r\n\r\n  planı\t\taktif  ".encode("utf-8"))
        )

        blocks = extract_document_text(document)

        self.assertEqual(blocks[0].content, "Üretim\n\nplanı aktif")

    def test_scanned_pdf_fails_without_ocr(self):
        document = self.create_document(blank_pdf_upload())

        result = process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertEqual(document.chunks.count(), 0)


class RejectedDocumentTests(ProcessingTestBase):
    def test_encrypted_pdf_is_rejected_during_upload_validation(self):
        with self.assertRaises(ValidationError):
            self.create_document(encrypted_pdf_upload())

    def test_macro_enabled_docx_is_rejected(self):
        with self.assertRaises(ValidationError):
            self.create_document(macro_docx_upload())

    def test_corrupt_and_unsupported_files_are_rejected(self):
        invalid_uploads = (
            SimpleUploadedFile("broken.pdf", b"%PDF-broken", content_type=PDF_MIME),
            SimpleUploadedFile("guide.exe", b"unsafe", content_type="application/octet-stream"),
        )
        for upload in invalid_uploads:
            with self.subTest(upload=upload.name), self.assertRaises(ValidationError):
                self.create_document(upload)

    def test_changed_stored_file_is_rejected_before_extraction(self):
        document = self.create_document()
        with document.file.storage.open(document.file.name, "wb") as stored_file:
            stored_file.write(b"changed after validation")

        result = process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(document.status, Document.Status.FAILED)


class ChunkingTests(TestCase):
    @staticmethod
    def long_text(word_count=2500):
        return " ".join(f"parca{index}" for index in range(word_count))

    def test_chunk_order_overlap_and_default_token_limits(self):
        chunks = create_chunks([TextBlock(content=self.long_text())])

        self.assertGreater(len(chunks), 2)
        self.assertEqual([chunk.chunk_index for chunk in chunks], list(range(len(chunks))))
        self.assertTrue(all(chunk.token_count <= 800 for chunk in chunks))
        self.assertTrue(all(600 <= chunk.token_count <= 800 for chunk in chunks[:-1]))

        first_words = chunks[0].content.split()
        second_words = chunks[1].content.split()
        shared_words = []
        for size in range(min(len(first_words), len(second_words)), 0, -1):
            if first_words[-size:] == second_words[:size]:
                shared_words = second_words[:size]
                break
        overlap_tokens = len(tiktoken.get_encoding("cl100k_base").encode(" ".join(shared_words)))
        self.assertGreaterEqual(overlap_tokens, 100)
        self.assertLessEqual(overlap_tokens, 150)

    @override_settings(
        AI_DOCUMENT_CHUNK_TARGET_TOKENS=40,
        AI_DOCUMENT_CHUNK_MAX_TOKENS=50,
        AI_DOCUMENT_CHUNK_OVERLAP_TOKENS=10,
        AI_DOCUMENT_EMBEDDING_INPUT_LIMIT_TOKENS=50,
    )
    def test_embedding_input_limit_is_a_hard_guard(self):
        chunks = create_chunks([TextBlock(content=self.long_text(200))])

        self.assertTrue(chunks)
        self.assertTrue(all(chunk.token_count <= 50 for chunk in chunks))

    def test_normalization_removes_controls_and_empty_blocks(self):
        self.assertEqual(normalize_text(" A\x00\t B\r\n\r\n\r\nC "), "A B\n\nC")
        chunks = create_chunks([TextBlock(content="..."), TextBlock(content="Anlamlı içerik")])
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].content, "Anlamlı içerik")
        self.assertEqual(
            chunks[0].content_hash,
            hashlib.sha256("Anlamlı içerik".encode("utf-8")).hexdigest(),
        )


class DocumentTaskTests(ProcessingTestBase):
    def test_new_queued_document_is_dispatched_after_commit(self):
        with patch("ai.signals.process_document.delay") as delay:
            with self.captureOnCommitCallbacks(execute=True):
                document = self.create_document()

        delay.assert_called_once_with(str(document.public_id))

    def test_success_sets_ready_and_creates_ordered_chunks(self):
        document = self.create_document()

        result = process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(result, {"status": "ready", "chunk_count": 1})
        self.assertEqual(document.status, Document.Status.READY)
        self.assertEqual(document.processing_error, "")
        self.assertEqual(list(document.chunks.values_list("chunk_index", flat=True)), [0])

    def test_temporary_storage_failure_is_requeued_for_limited_retry(self):
        document = self.create_document()

        with patch("ai.tasks.build_document_chunks", side_effect=OSError("temporary")):
            with self.assertRaises(OSError):
                process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUEUED)
        self.assertIn("OSError", document.processing_error)
        self.assertEqual(process_document.max_retries, 2)

    def test_exhausted_temporary_retries_set_failed(self):
        document = self.create_document()
        process_document.push_request(retries=2, called_directly=False)
        try:
            with patch(
                "ai.tasks.build_document_chunks",
                side_effect=OSError("still unavailable"),
            ):
                result = process_document.run(str(document.public_id))
        finally:
            process_document.pop_request()

        document.refresh_from_db()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertIn("OSError", document.processing_error)

    def test_reprocessing_atomically_replaces_old_chunks_and_is_idempotent(self):
        document = self.create_document()
        process_document.run(str(document.public_id))
        original_hashes = list(document.chunks.values_list("content_hash", flat=True))
        DocumentChunk.objects.create(
            document=document,
            chunk_index=99,
            content="stale",
            token_count=1,
            content_hash=hashlib.sha256(b"stale").hexdigest(),
        )

        Document.objects.filter(pk=document.pk).update(status=Document.Status.QUEUED)
        result = process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(result["status"], "ready")
        self.assertEqual(document.status, Document.Status.READY)
        self.assertEqual(
            list(document.chunks.values_list("content_hash", flat=True)),
            original_hashes,
        )
        self.assertEqual(document.chunks.count(), len(set(original_hashes)))

    def test_failure_keeps_old_chunks_and_sets_failed(self):
        document = self.create_document()
        process_document.run(str(document.public_id))
        old_hashes = list(document.chunks.values_list("content_hash", flat=True))
        Document.objects.filter(pk=document.pk).update(status=Document.Status.QUEUED)

        with patch("ai.tasks.build_document_chunks", side_effect=ValueError("technical detail")):
            result = process_document.run(str(document.public_id))

        document.refresh_from_db()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertIn("ValueError", document.processing_error)
        self.assertEqual(
            list(document.chunks.values_list("content_hash", flat=True)),
            old_hashes,
        )

    def test_deleting_document_cascades_chunks_and_removes_private_file(self):
        document = self.create_document()
        process_document.run(str(document.public_id))
        file_name = document.file.name
        storage = document.file.storage
        self.assertTrue(storage.exists(file_name))
        self.assertTrue(document.chunks.exists())

        document.delete()

        self.assertFalse(DocumentChunk.objects.exists())
        self.assertFalse(storage.exists(file_name))

    def test_buyer_cannot_use_reprocessing_admin_action(self):
        buyer = get_user_model().objects.create(username="processing-buyer", is_staff=True)
        Group.objects.create(name="Buyer").user_set.add(buyer)
        document = self.create_document()
        Document.objects.filter(pk=document.pk).update(status=Document.Status.READY)
        self.client.force_login(buyer)

        response = self.client.post(
            reverse("admin:ai_document_changelist"),
            {
                "action": "queue_for_reprocessing",
                "_selected_action": [document.pk],
            },
        )

        self.assertEqual(response.status_code, 403)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.READY)

    def test_staff_can_queue_reprocessing_after_commit(self):
        document = self.create_document()
        Document.objects.filter(pk=document.pk).update(status=Document.Status.READY)
        self.client.force_login(self.staff)

        with patch("ai.signals.process_document.delay") as delay:
            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post(
                    reverse("admin:ai_document_changelist"),
                    {
                        "action": "queue_for_reprocessing",
                        "_selected_action": [document.pk],
                    },
                )

        self.assertEqual(response.status_code, 302)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUEUED)
        delay.assert_called_once_with(str(document.public_id))


class ConcurrentDocumentTaskTests(TransactionTestCase):
    reset_sequences = True
    available_apps = ["django.contrib.auth", "django.contrib.contenttypes", "ai"]

    def setUp(self):
        super().setUp()
        self.media_directory = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.media_directory.name)
        self.settings_override.enable()
        self.staff = get_user_model().objects.create(username="concurrent-staff", is_staff=True)
        with patch("ai.signals.process_document.delay"):
            self.document = Document.objects.create(
                title="Concurrent",
                file=txt_upload(b"Es zamanli isleme testi icin guvenli metin"),
                uploaded_by=self.staff,
            )

    def tearDown(self):
        self.settings_override.disable()
        self.media_directory.cleanup()
        super().tearDown()

    def test_two_workers_do_not_create_duplicate_chunks(self):
        started = threading.Event()
        release = threading.Event()
        real_builder = process_document.run.__globals__["build_document_chunks"]
        build_calls = []

        def blocking_builder(document):
            build_calls.append(str(document.public_id))
            started.set()
            release.wait(timeout=10)
            return real_builder(document)

        def run_task():
            close_old_connections()
            try:
                return process_document.run(str(self.document.public_id))
            finally:
                close_old_connections()

        with patch("ai.tasks.build_document_chunks", side_effect=blocking_builder):
            with patch("ai.tasks.dispatch_document_embeddings"), ThreadPoolExecutor(max_workers=2) as executor:
                first = executor.submit(run_task)
                self.assertTrue(started.wait(timeout=10))
                second = executor.submit(run_task)
                release.set()
                second_result = second.result(timeout=10)
                first_result = first.result(timeout=10)

        self.document.refresh_from_db()
        self.assertEqual(first_result["status"], "ready")
        self.assertEqual(second_result["status"], "skipped")
        self.assertEqual(len(build_calls), 1)
        self.assertEqual(self.document.status, Document.Status.READY)
        self.assertEqual(self.document.chunks.count(), 1)
        self.assertEqual(
            self.document.chunks.values("chunk_index").distinct().count(),
            self.document.chunks.count(),
        )
        self.assertEqual(
            self.document.chunks.values("content_hash").distinct().count(),
            self.document.chunks.count(),
        )
