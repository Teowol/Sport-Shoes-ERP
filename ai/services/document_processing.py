"""Safe text extraction, normalization, and token chunking for AI documents."""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from bisect import bisect_right
from dataclasses import dataclass
from pathlib import PurePosixPath

import tiktoken
from django.conf import settings
from django.core.exceptions import ValidationError
from docx import Document as DocxDocument
from pypdf import PdfReader

from ai.document_validation import ALLOWED_DOCUMENT_TYPES, validate_docx_archive


class DocumentProcessingError(Exception):
    """A permanent, expected rejection while processing a document."""


@dataclass(frozen=True)
class TextBlock:
    content: str
    page_number: int | None = None
    section_title: str = ""


@dataclass(frozen=True)
class ChunkData:
    chunk_index: int
    content: str
    token_count: int
    page_number: int | None
    section_title: str
    content_hash: str


def normalize_text(value: str) -> str:
    """Normalize Unicode and whitespace without flattening paragraph boundaries."""
    text = unicodedata.normalize("NFKC", value or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = "".join(
        character
        for character in text
        if character in "\n\t" or not unicodedata.category(character).startswith("C")
    )
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def is_meaningful_text(value: str) -> bool:
    return sum(character.isalnum() for character in value) >= 3


def _read_verified_payload(document) -> tuple[bytes, str]:
    extension = PurePosixPath(document.original_filename).suffix.lower()
    if extension not in ALLOWED_DOCUMENT_TYPES:
        raise DocumentProcessingError("Desteklenmeyen doküman türü.")
    if document.mime_type not in ALLOWED_DOCUMENT_TYPES[extension]:
        raise DocumentProcessingError("Doküman türü ve MIME bilgisi eşleşmiyor.")

    max_size = int(getattr(settings, "AI_DOCUMENT_MAX_SIZE_BYTES", 10 * 1024 * 1024))
    try:
        with document.file.storage.open(document.file.name, "rb") as stored_file:
            payload = stored_file.read(max_size + 1)
    except (OSError, ValueError) as exc:
        raise OSError("Doküman depolama alanından okunamadı.") from exc

    if not payload or len(payload) > max_size:
        raise DocumentProcessingError("Doküman boyutu güvenli sınırlar dışında.")
    if len(payload) != document.file_size:
        raise DocumentProcessingError("Saklanan dokümanın boyutu değişmiş.")
    if hashlib.sha256(payload).hexdigest() != document.checksum_sha256:
        raise DocumentProcessingError("Saklanan dokümanın checksum değeri değişmiş.")
    return payload, extension


def _extract_pdf(payload: bytes) -> list[TextBlock]:
    if not payload.startswith(b"%PDF-"):
        raise DocumentProcessingError("PDF dosya imzası geçersiz.")
    try:
        reader = PdfReader(io.BytesIO(payload), strict=True)
        if reader.is_encrypted:
            raise DocumentProcessingError("Şifreli PDF dosyaları desteklenmiyor.")
        blocks = [
            TextBlock(content=normalize_text(page.extract_text() or ""), page_number=index)
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("PDF metni çıkarılamadı.") from exc
    if not any(is_meaningful_text(block.content) for block in blocks):
        raise DocumentProcessingError("Taranmış PDF ve OCR bu aşamada desteklenmiyor.")
    return blocks


def _extract_docx(payload: bytes) -> list[TextBlock]:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            validate_docx_archive(archive, len(payload))
        document = DocxDocument(io.BytesIO(payload))
    except ValidationError as exc:
        raise DocumentProcessingError(str(exc)) from exc
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise DocumentProcessingError("DOCX metni çıkarılamadı.") from exc

    blocks: list[TextBlock] = []
    section_title = ""
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not is_meaningful_text(text):
            continue
        style_name = normalize_text(getattr(paragraph.style, "name", ""))
        if style_name.casefold().startswith("heading"):
            section_title = text[:255]
            continue
        blocks.append(TextBlock(content=text, section_title=section_title))

    for table in document.tables:
        for row in table.rows:
            text = normalize_text(" | ".join(cell.text for cell in row.cells))
            if is_meaningful_text(text):
                blocks.append(TextBlock(content=text, section_title=section_title))

    if not blocks:
        raise DocumentProcessingError("DOCX belgesinde anlamlı metin bulunamadı.")
    return blocks


def _extract_txt(payload: bytes) -> list[TextBlock]:
    if b"\x00" in payload:
        raise DocumentProcessingError("TXT dosyası ikili veri içeremez.")
    try:
        text = normalize_text(payload.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("TXT dosyası geçerli UTF-8 değil.") from exc
    if not is_meaningful_text(text):
        raise DocumentProcessingError("TXT belgesinde anlamlı metin bulunamadı.")
    return [TextBlock(content=text)]


def extract_document_text(document) -> list[TextBlock]:
    payload, extension = _read_verified_payload(document)
    if extension == ".pdf":
        return _extract_pdf(payload)
    if extension == ".docx":
        return _extract_docx(payload)
    if extension == ".txt":
        return _extract_txt(payload)
    raise DocumentProcessingError("Desteklenmeyen doküman türü.")


def _chunk_settings() -> tuple[int, int, int, int, str]:
    target = int(getattr(settings, "AI_DOCUMENT_CHUNK_TARGET_TOKENS", 700))
    maximum = int(getattr(settings, "AI_DOCUMENT_CHUNK_MAX_TOKENS", 800))
    overlap = int(getattr(settings, "AI_DOCUMENT_CHUNK_OVERLAP_TOKENS", 120))
    embedding_limit = int(
        getattr(settings, "AI_DOCUMENT_EMBEDDING_INPUT_LIMIT_TOKENS", 8192)
    )
    encoding_name = str(getattr(settings, "AI_DOCUMENT_TOKEN_ENCODING", "cl100k_base"))
    if not (0 < overlap < target <= maximum <= embedding_limit):
        raise DocumentProcessingError("Doküman chunk token ayarları geçersiz.")
    return target, maximum, overlap, embedding_limit, encoding_name


def _starts_at_word_boundary(encoding, token: int) -> bool:
    decoded = encoding.decode([token])
    return bool(decoded) and decoded[0].isspace()


def _align_chunk_end(encoding, token_stream: list[int], start: int, end: int) -> int:
    """Avoid ending inside a word while keeping the chunk close to its target."""
    lower_bound = max(start + 1, end - 20)
    for candidate in range(end, lower_bound - 1, -1):
        if candidate < len(token_stream) and _starts_at_word_boundary(
            encoding, token_stream[candidate]
        ):
            return candidate
    return end


def _align_overlap_start(
    encoding, token_stream: list[int], start: int, end: int
) -> int:
    """Keep overlap near 120 tokens and start the next chunk on a whole word."""
    for candidate in range(start, min(end, start + 20) + 1):
        if _starts_at_word_boundary(encoding, token_stream[candidate]):
            return candidate
    return start


def create_chunks(blocks: list[TextBlock]) -> list[ChunkData]:
    target, maximum, overlap, embedding_limit, encoding_name = _chunk_settings()
    try:
        encoding = tiktoken.get_encoding(encoding_name)
    except Exception as exc:
        raise DocumentProcessingError("Token kodlama ayarı yüklenemedi.") from exc

    token_stream: list[int] = []
    span_starts: list[int] = []
    span_metadata: list[tuple[int | None, str]] = []
    separator = encoding.encode("\n\n")
    for block in blocks:
        content = normalize_text(block.content)
        if not is_meaningful_text(content):
            continue
        if token_stream:
            token_stream.extend(separator)
        span_starts.append(len(token_stream))
        span_metadata.append((block.page_number, block.section_title[:255]))
        token_stream.extend(encoding.encode(content))

    if not token_stream:
        raise DocumentProcessingError("Dokümanda anlamlı metin bulunamadı.")

    chunks: list[ChunkData] = []
    seen_hashes: set[str] = set()
    start = 0
    while start < len(token_stream):
        remaining = len(token_stream) - start
        window_size = min(remaining, maximum if remaining <= maximum else target)
        end = start + window_size
        if end < len(token_stream):
            end = _align_chunk_end(encoding, token_stream, start, end)
        content = normalize_text(encoding.decode(token_stream[start:end]))
        actual_tokens = len(encoding.encode(content))
        if is_meaningful_text(content) and actual_tokens:
            if actual_tokens > maximum or actual_tokens > embedding_limit:
                raise DocumentProcessingError("Oluşturulan chunk token sınırını aşıyor.")
            content_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()
            if content_hash not in seen_hashes:
                metadata_index = max(0, bisect_right(span_starts, start) - 1)
                page_number, section_title = span_metadata[metadata_index]
                chunks.append(
                    ChunkData(
                        chunk_index=len(chunks),
                        content=content,
                        token_count=actual_tokens,
                        page_number=page_number,
                        section_title=section_title,
                        content_hash=content_hash,
                    )
                )
                seen_hashes.add(content_hash)
        if end >= len(token_stream):
            break
        start = _align_overlap_start(
            encoding, token_stream, max(0, end - overlap), end
        )

    if not chunks:
        raise DocumentProcessingError("Dokümandan kullanılabilir chunk oluşturulamadı.")
    return chunks


def build_document_chunks(document) -> list[ChunkData]:
    return create_chunks(extract_document_text(document))
